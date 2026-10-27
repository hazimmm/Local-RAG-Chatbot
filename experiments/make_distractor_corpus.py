"""Distractor documents: same vocabulary as the main corpus, different topics.

They must NOT answer any golden-set question; they only compete for rank.
"""
from pathlib import Path

from docx import Document

OUT = Path(__file__).resolve().parent.parent / "samples" / "distractors"
OUT.mkdir(parents=True, exist_ok=True)

PROCUREMENT = [
    ("h1", "Acme Bank Procurement Policy"),
    ("p", "This policy governs purchases of goods and services from "
          "suppliers. It does not cover employee expense claims."),
    ("h2", "Purchase Order Approval Limits"),
    ("table", [["Order value", "Approver"],
               ["Up to 500 dollars", "Team lead"],
               ["501 to 5,000 dollars", "Procurement manager"],
               ["Over 5,000 dollars", "Chief operating officer"]]),
    ("h2", "Supplier Invoices"),
    ("p", "Approved supplier invoices are paid within 45 days of receipt. "
          "Invoices received after the 25th are processed in the following "
          "month's payment run."),
    ("h2", "Supplier Travel"),
    ("p", "Suppliers visiting Acme Bank sites are not reimbursed for hotels "
          "or meals unless the contract says so. Contracted hotel stays are "
          "capped at 140 dollars per night."),
    ("h2", "Equipment Purchases"),
    ("p", "Laptops and phones are ordered through the central equipment "
          "catalogue. Replacement of damaged equipment needs a written "
          "incident report within 5 working days."),
]

FACILITIES = [
    ("h1", "Acme Bank Facilities and Security Guide"),
    ("p", "This guide explains access to Acme Bank buildings and shared "
          "facilities."),
    ("h2", "Access Badges"),
    ("p", "Every visitor and employee must wear a badge at all times. Lost "
          "or stolen badges must be reported to reception within 1 hour. "
          "A replacement badge costs 15 dollars."),
    ("h2", "Guest Wi-Fi"),
    ("p", "Visitors may use the guest Wi-Fi network. The guest password "
          "changes every 7 days and is available at reception. Guest "
          "devices cannot reach internal systems."),
    ("h2", "Canteen and Meals"),
    ("p", "The staff canteen is open from 8am to 3pm Monday to Friday. A "
          "lunch voucher worth 6 dollars is provided to all staff on "
          "working days."),
    ("h2", "Building Hours"),
    ("p", "Head office is open to staff from 7am to 8pm Monday to Friday "
          "and is closed at weekends. Out-of-hours access needs approval "
          "from the facilities manager."),
    ("h2", "Emergency Contacts"),
    ("table", [["Situation", "Contact"],
               ["Fire or medical emergency", "Reception, extension 4000"],
               ["Building faults", "Facilities desk, extension 4400"],
               ["Security incidents", "Security office, extension 4500"]]),
]

INSURANCE_FAQ = """# Acme Insurance Customer FAQ

## Claims

### How do I make a claim?
Submit your claim online within 30 days of the incident with original
receipts. Claims over 1,000 dollars are reviewed by a senior assessor.

### How long does a claim take to pay?
Approved claims are paid within 10 working days. Claims submitted after the
20th of the month may take an extra week.

## Cards and Policies

### What should I do if my insurance card is lost?
Request a new insurance card in the mobile app, or call 0800 555 0142.
A replacement card arrives within 7 working days and costs nothing.

### Can I transfer my policy to another person?
Policies can be transferred to a family member once per year. A transfer
fee of 20 dollars applies.

## Premiums

### What is the premium for travel insurance?
Single trip travel insurance costs 35 dollars. Annual multi-trip cover
costs 120 dollars per year and includes hotel cancellation cover up to 200
dollars per night.

## Office Hours

### When are the insurance service centres open?
Service centres are open Monday to Thursday from 9am to 4pm and on Fridays
from 9am to 12pm. They are closed at weekends.
"""


def write_docx(name: str, blocks: list) -> None:
    doc = Document()
    for kind, content in blocks:
        if kind == "h1":
            doc.add_heading(content, level=1)
        elif kind == "h2":
            doc.add_heading(content, level=2)
        elif kind == "p":
            doc.add_paragraph(content)
        elif kind == "table":
            table = doc.add_table(rows=len(content), cols=len(content[0]))
            table.style = "Table Grid"
            for r, row in enumerate(content):
                for c, text in enumerate(row):
                    table.rows[r].cells[c].text = text
    doc.save(OUT / name)


write_docx("procurement_policy.docx", PROCUREMENT)
write_docx("facilities_guide.docx", FACILITIES)
(OUT / "insurance_faq.md").write_text(INSURANCE_FAQ, encoding="utf-8")
print("written:", sorted(p.name for p in OUT.iterdir()))
