"""Generate a small fictional corpus (Acme Bank) in mixed file formats."""
from pathlib import Path

from docx import Document
from fpdf import FPDF
from openpyxl import Workbook

OUT = Path(__file__).resolve().parent.parent / "samples" / "corpus"
OUT.mkdir(parents=True, exist_ok=True)

# Each document is a list of blocks:
# ("h1", text) ("h2", text) ("p", text) ("ul", [items]) ("table", rows)

HR = [
    ("h1", "Acme Bank Leave and Benefits Policy"),
    ("p", "This policy applies to all permanent employees of Acme Bank. "
          "Contract staff should refer to their agreement."),
    ("h2", "Annual Leave"),
    ("p", "Full-time employees receive 20 days of annual leave per calendar "
          "year. Leave is requested through the HR portal at least 14 days "
          "in advance and must be approved by the line manager."),
    ("p", "Up to 5 unused days can be carried over to the next year. Carried "
          "over days expire on 31 March."),
    ("h2", "Sick Leave"),
    ("p", "Employees receive 10 days of paid sick leave per year. A medical "
          "certificate is required when sick leave lasts more than 3 "
          "consecutive days. Employees must notify their manager before 9am "
          "on the first day of absence."),
    ("h2", "Parental Leave"),
    ("p", "Primary caregivers receive 14 weeks of paid parental leave. "
          "Secondary caregivers receive 4 weeks. Leave may begin up to 4 "
          "weeks before the expected due date."),
    ("h2", "Bereavement and Unpaid Leave"),
    ("ul", ["Bereavement leave: 5 days for immediate family, 2 days for "
            "extended family.",
            "Unpaid leave of up to 30 days may be granted after 2 years of "
            "service, subject to director approval."]),
    ("h2", "Health Benefits"),
    ("p", "All employees are enrolled in the company medical insurance plan "
          "from their first day. Dependants can be added during the open "
          "enrolment period every November. Dental cover is optional and "
          "costs 25 dollars per month."),
    ("h2", "Probation"),
    ("p", "The probation period is 3 months. Annual leave accrues during "
          "probation but may only be taken after it ends, except with "
          "manager approval."),
]

IT = [
    ("h1", "Acme Bank Information Security Policy"),
    ("p", "All staff must follow this policy when using Acme Bank systems "
          "and data."),
    ("h2", "Passwords and Authentication"),
    ("p", "Passwords must be at least 12 characters long and include "
          "letters, numbers and a symbol. Passwords expire every 90 days "
          "and the last 5 passwords cannot be reused."),
    ("p", "Multi-factor authentication (MFA) is mandatory for email, the VPN "
          "and the core banking system. To reset a forgotten password, use "
          "the IT portal or call the service desk on extension 4100."),
    ("h2", "Devices"),
    ("ul", ["Company laptops must use full-disk encryption.",
            "Personal devices may access email only through the managed "
            "mobile app.",
            "Lost or stolen devices must be reported to the service desk "
            "within 2 hours."]),
    ("h2", "Remote Access"),
    ("p", "Remote work requires the company VPN. Public Wi-Fi may be used "
          "only with the VPN active. Customer data must never be saved on "
          "personal cloud storage."),
    ("h2", "Phishing and Incident Reporting"),
    ("p", "Suspicious emails must be reported with the Report Phish button. "
          "Any suspected data breach must be reported to the security team "
          "at once. The security team will classify the incident within 24 "
          "hours."),
    ("h2", "Data Classification"),
    ("table", [["Level", "Example", "Handling"],
               ["Public", "Marketing brochures", "No restrictions"],
               ["Internal", "Staff directory", "Company staff only"],
               ["Confidential", "Customer account records",
                "Encrypted, need-to-know access"]]),
]

EXPENSE = [
    ("h1", "Acme Bank Travel and Expense Policy"),
    ("p", "Employees are reimbursed for reasonable business expenses. All "
          "claims must be submitted within 30 days with original receipts."),
    ("h2", "Approval Limits"),
    ("table", [["Amount", "Approver"],
               ["Up to 200 dollars", "Line manager"],
               ["201 to 2,000 dollars", "Department head"],
               ["Over 2,000 dollars", "Finance director"]]),
    ("h2", "Meals"),
    ("p", "Daily meal allowance during business travel is 60 dollars. "
          "Alcohol is not reimbursed under any circumstances."),
    ("h2", "Flights and Hotels"),
    ("p", "Flights under 6 hours must be booked in economy class. Hotel "
          "stays are capped at 180 dollars per night in major cities and "
          "120 dollars elsewhere. Bookings must be made through the "
          "company travel portal."),
    ("h2", "Reimbursement Timeline"),
    ("p", "Approved claims are paid with the next monthly payroll. Claims "
          "submitted after the 20th are paid the following month."),
    ("h2", "Mileage"),
    ("p", "Business use of a personal car is reimbursed at 0.58 dollars per "
          "kilometre."),
]

FAQ_MD = """# Acme Bank Customer FAQ

## Accounts

### How do I open a savings account?
You can open a savings account online or at any branch. You need a valid
photo ID and proof of address. The minimum opening deposit is 100 dollars.

### What is the interest rate on savings accounts?
The Everyday Saver pays 2.1 percent per year and the Premium Saver pays
3.4 percent per year for balances above 10,000 dollars.

## Cards

### What should I do if my card is lost or stolen?
Freeze the card in the mobile app straight away, then call 0800 555 0199
at any time. A replacement card arrives within 5 working days.

### What are the daily withdrawal limits?
The standard debit card limit is 1,000 dollars per day at ATMs.

## Transfers

### How long do transfers take?
Transfers between Acme Bank accounts are instant. Domestic transfers to
other banks take 1 business day. International transfers take 2 to 4
business days.

### What are the transfer limits?
Online transfers are limited to 25,000 dollars per day. Higher amounts
require a branch visit.

## Branch Services

### What are the branch opening hours?
Branches are open Monday to Friday from 9am to 5pm and on Saturdays from
10am to 1pm. Branches are closed on Sundays and public holidays.
"""

ONBOARDING = [
    ("h1", "New Employee Onboarding Guide"),
    ("p", "Welcome to Acme Bank. This guide covers your first week."),
    ("h2", "Day One"),
    ("p", "Report to reception at 9am. You will receive your ID badge and "
          "laptop. IT will help you set up your account and multi-factor "
          "authentication."),
    ("h2", "First Week Schedule"),
    ("ul", ["Monday: welcome session and compliance training.",
            "Tuesday: meet your team and manager.",
            "Wednesday: systems training.",
            "Thursday: shadow a colleague.",
            "Friday: check-in with HR."]),
    ("h2", "Mandatory Training"),
    ("p", "All new hires complete anti-money-laundering and data privacy "
          "training within 30 days. Failure to complete training may delay "
          "system access."),
    ("h2", "Key Contacts"),
    ("table", [["Topic", "Contact"],
               ["IT problems", "Service desk, extension 4100"],
               ["Leave and payroll", "HR team, extension 4200"],
               ["Expenses", "Finance team, extension 4300"]]),
    ("h2", "Working Hours"),
    ("p", "Head office hours are 8:30am to 5:30pm Monday to Friday, with a "
          "one-hour lunch break. Flexible start times between 7:30am and "
          "9:30am can be agreed with your manager."),
]

FEES = {
    "Account fees": [
        ["Account type", "Monthly fee", "Minimum balance", "Free ATM withdrawals"],
        ["Basic Current", "0", "0", "4"],
        ["Plus Current", "8", "1000", "10"],
        ["Everyday Saver", "0", "100", "2"],
        ["Premium Saver", "5", "10000", "Unlimited"],
    ],
    "Service charges": [
        ["Service", "Fee in dollars"],
        ["Replacement card", "10"],
        ["International transfer", "15"],
        ["Cheque book", "12"],
        ["Late payment", "25"],
    ],
}


def write_docx(name: str, blocks: list) -> None:
    doc = Document()
    for kind, content in blocks:
        if kind == "h1":
            doc.add_heading(content, level=1)
        elif kind == "h2":
            doc.add_heading(content, level=2)
        elif kind == "p":
            doc.add_paragraph(content)
        elif kind == "ul":
            for item in content:
                doc.add_paragraph(item, style="List Bullet")
        elif kind == "table":
            table = doc.add_table(rows=len(content), cols=len(content[0]))
            table.style = "Table Grid"
            for r, row in enumerate(content):
                for c, text in enumerate(row):
                    table.rows[r].cells[c].text = text
    doc.save(OUT / name)


def write_pdf(name: str, blocks: list) -> None:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    for kind, content in blocks:
        pdf.set_x(pdf.l_margin)
        if kind == "h1":
            pdf.set_font("Helvetica", "B", 18)
            pdf.multi_cell(0, 10, content)
        elif kind == "h2":
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 13)
            pdf.multi_cell(0, 8, content)
        elif kind == "p":
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(0, 6, content)
        elif kind == "ul":
            pdf.set_font("Helvetica", "", 11)
            for item in content:
                pdf.set_x(pdf.l_margin)
                pdf.multi_cell(0, 6, "- " + item)
        elif kind == "table":
            pdf.set_font("Helvetica", "", 11)
            for row in content:
                pdf.set_x(pdf.l_margin)
                pdf.multi_cell(0, 6, "  |  ".join(row))
    pdf.output(str(OUT / name))


def write_xlsx(name: str, sheets: dict) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append([int(v) if v.isdigit() else v for v in row])
    wb.save(OUT / name)


write_docx("hr_leave_and_benefits.docx", HR)
write_docx("it_security_policy.docx", IT)
write_docx("expense_policy.docx", EXPENSE)
(OUT / "customer_faq.md").write_text(FAQ_MD, encoding="utf-8")
write_pdf("onboarding_guide.pdf", ONBOARDING)
write_xlsx("fee_schedule.xlsx", FEES)
print("written:", sorted(p.name for p in OUT.iterdir()))
