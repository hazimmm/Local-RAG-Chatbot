from pathlib import Path

from docx import Document

doc = Document()
doc.add_heading("Acme Bank Employee Handbook", level=1)
doc.add_paragraph("This handbook explains common policies for Acme Bank staff.")

doc.add_heading("Password Policy", level=2)
doc.add_paragraph("Passwords must be at least 12 characters long and changed every 90 days.")
doc.add_paragraph("To reset a forgotten password, open the IT portal and choose Reset Credentials.")

doc.add_heading("Branch Hours", level=2)
table = doc.add_table(rows=3, cols=2)
table.style = "Table Grid"
rows = [("Day", "Hours"), ("Monday to Friday", "9am to 5pm"), ("Saturday", "10am to 1pm")]
for row, (day, hours) in zip(table.rows, rows):
    row.cells[0].text = day
    row.cells[1].text = hours

doc.add_heading("Leave Policy", level=2)
for item in ["Annual leave: 20 days", "Sick leave: 10 days", "Parental leave: 14 weeks"]:
    doc.add_paragraph(item, style="List Bullet")

out = Path(__file__).resolve().parent.parent / "samples" / "handbook.docx"
doc.save(out)
print("saved", out)
